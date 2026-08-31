"""
populate_catalogos.py — Seeder Idempotente de Catálogos de Referencia 3NF

Lee exclusivamente los datos reales de uploads/data/ (no tests ni sintéticos):
- 'Base centralizada Académica ARCA (1).xlsx' (CRITERIOS, IBP, MODULOS)
- 'EJECUTADO VS META.xlsx' (Metas 2024-2026)
- Definición canónica de carteras de cobranza (Dashboard Cobranzas)

Extrae:
- Escuelas, tipos de programa, modalidades
- Programas (parseando versión numérica ej. v. 28, v27, etc.)
- Docentes canónicos (nombre, email, celular, grado académico, tipo)
- Módulos con horas y pagos asignados
- Criterios académicos (estados académicos, reprobados, ARCA)
- Clasificación de mora/carteras
- Metas presupuestarias mensuales

Genera 'db/seeds/seed_catalogos.sql' con ON CONFLICT para idempotencia
e intenta insertar en PostgreSQL (giragroup_catalog_db) si está accesible.
"""

import os
import re
import sys
import openpyxl
from pathlib import Path


def limpiar_texto(val) -> str:
    if val is None:
        return ""
    txt = str(val).strip()
    # Eliminar dobles espacios
    txt = re.sub(r'\s+', ' ', txt)
    return txt


def escapar_sql(val) -> str:
    if val is None:
        return "NULL"
    s = str(val).replace("'", "''")
    return f"'{s}'"


def extraer_datos_catalogos(dir_datos: str = "uploads/data") -> dict:
    catalogos = {
        "escuelas": set(),
        "tipos_programa": set(),
        "modalidades": set(),
        "programas": {},      # pos_code -> dict
        "docentes": {},       # nombre_completo -> dict
        "modulos": [],        # list of dicts
        "estados_academicos": set(),
        "estados_reprobados": set(),
        "estados_arca": set(),
        "carteras": [
            {"nombre": "Cartera Vigente", "min": 0, "max": 90, "desc": "Gestión activa y controlada (0-90 días)"},
            {"nombre": "Cartera en Mora", "min": 91, "max": 180, "desc": "Retraso intermedio (91-180 días)"},
            {"nombre": "Cartera Previsionable", "min": 181, "max": 360, "desc": "Situación crítica en seguimiento (181-360 días)"},
            {"nombre": "Cartera Incobrable", "min": 361, "max": 99999, "desc": "Sin posibilidades operativas (>360 días)"},
            {"nombre": "Cartera Cerrada", "min": -1, "max": -1, "desc": "Caso cerrado satisfactoriamente (sin pendientes)"},
        ],
        "metas": [],          # list of dicts
    }

    # ──────────────────────────────────────────────────────────────────────────
    # 1. Base centralizada Académica ARCA (1).xlsx
    # ──────────────────────────────────────────────────────────────────────────
    ruta_arca = os.path.join(dir_datos, "Base centralizada Académica ARCA (1).xlsx")
    if os.path.exists(ruta_arca):
        print(f"Leyendo '{ruta_arca}'...")
        wb = openpyxl.load_workbook(ruta_arca, data_only=True)

        # A. Criterios Académicos
        if "CRITERIOS" in wb.sheetnames:
            ws = wb["CRITERIOS"]
            for row in list(ws.iter_rows(values_only=True))[1:]:
                # col 1: Estado academico, col 2: Estado reprobado, col 3: Estado ARCA
                if len(row) > 1 and row[1]:
                    v = limpiar_texto(row[1])
                    if v: catalogos["estados_academicos"].add(v)
                if len(row) > 2 and row[2]:
                    v = limpiar_texto(row[2])
                    if v: catalogos["estados_reprobados"].add(v)
                if len(row) > 3 and row[3]:
                    v = limpiar_texto(row[3])
                    if v: catalogos["estados_arca"].add(v)

        # B. IBP (Programas, Escuelas, Modalidades, Tipos)
        if "IBP" in wb.sheetnames:
            ws = wb["IBP"]
            for row in list(ws.iter_rows(values_only=True))[1:]:
                if len(row) < 13:
                    continue
                pos_raw = row[1]
                nombre_raw = row[2]
                ver_col = row[3]
                escuela_raw = row[6]
                inversion = row[7]
                punto_eq = row[8]
                horas = row[9]
                modalidad_raw = row[11]
                tipo_raw = row[12]

                if not pos_raw or not nombre_raw:
                    continue

                pos_clean = limpiar_texto(pos_raw)
                if pos_clean.replace(".", "").isdigit():
                    pos_code = f"POS-{int(float(pos_clean))}"
                else:
                    pos_code = pos_clean if pos_clean.startswith("POS-") else f"POS-{pos_clean}"

                # Escuelas, Tipos, Modalidades
                escuela = limpiar_texto(escuela_raw)
                if escuela:
                    catalogos["escuelas"].add(escuela)

                tipo = limpiar_texto(tipo_raw).upper()
                if tipo:
                    catalogos["tipos_programa"].add(tipo)

                modalidad = limpiar_texto(modalidad_raw).upper()
                if modalidad:
                    catalogos["modalidades"].add(modalidad)

                # Detección de versión en el nombre (ej. v. 28, v27, v.14) o en columna VERSIÓN
                nombre_str = limpiar_texto(nombre_raw)
                match_ver = re.search(r'\bv\.?\s*(\d+)', nombre_str, re.IGNORECASE)
                if match_ver:
                    version = int(match_ver.group(1))
                elif ver_col and str(ver_col).replace(".", "").isdigit():
                    version = int(float(ver_col))
                else:
                    version = 1

                # Nombre canónico limpio (sin sufijo de versión)
                nombre_canonico = re.sub(r'\bv\.?\s*\d+', '', nombre_str, flags=re.IGNORECASE).strip()
                nombre_canonico = re.sub(r'\s+', ' ', nombre_canonico)

                inv_val = float(inversion) if inversion and str(inversion).replace(".", "").isdigit() else 0.0
                peq_val = int(float(punto_eq)) if punto_eq and str(punto_eq).replace(".", "").isdigit() else 0
                hrs_val = int(float(horas)) if horas and str(horas).replace(".", "").isdigit() else 0

                catalogos["programas"][pos_code] = {
                    "pos_code": pos_code,
                    "nombre_programa": nombre_canonico,
                    "version": version,
                    "nombre_completo_raw": nombre_str,
                    "escuela": escuela,
                    "tipo_programa": tipo,
                    "modalidad": modalidad,
                    "inversion_base": inv_val,
                    "punto_equilibrio": peq_val,
                    "horas_totales": hrs_val,
                }

        # C. MODULOS y DOCENTES
        if "MODULOS" in wb.sheetnames:
            ws = wb["MODULOS"]
            for row in list(ws.iter_rows(values_only=True))[1:]:
                if len(row) < 15:
                    continue
                pos_raw = row[1]
                doc_nombre = row[6]
                tipo_doc = row[7]
                nro_mod = row[8]
                nombre_mod = row[9]
                hrs_mod = row[10]
                pago_doc = row[11]
                grado = row[12]
                cel = row[13]
                email = row[14]

                # Docente
                if doc_nombre:
                    d_clean = limpiar_texto(doc_nombre)
                    if d_clean and d_clean not in catalogos["docentes"]:
                        catalogos["docentes"][d_clean] = {
                            "nombre_completo": d_clean,
                            "tipo_docente": limpiar_texto(tipo_doc) if tipo_doc else "TITULAR",
                            "grado_academico": limpiar_texto(grado) if grado else None,
                            "celular": limpiar_texto(cel) if cel else None,
                            "email": limpiar_texto(email) if email else None,
                        }

                # Módulo
                if pos_raw and nro_mod and nombre_mod:
                    pos_clean = limpiar_texto(pos_raw)
                    if pos_clean.replace(".", "").isdigit():
                        pos_code = f"POS-{int(float(pos_clean))}"
                    else:
                        pos_code = pos_clean if pos_clean.startswith("POS-") else f"POS-{pos_clean}"

                    try:
                        nro_val = int(float(nro_mod))
                    except (ValueError, TypeError):
                        continue

                    try:
                        hrs_val = float(hrs_mod) if hrs_mod else 0.0
                    except (ValueError, TypeError):
                        hrs_val = 0.0

                    try:
                        pago_val = float(pago_doc) if pago_doc else 0.0
                    except (ValueError, TypeError):
                        pago_val = 0.0

                    catalogos["modulos"].append({
                        "pos_code": pos_code,
                        "nro_modulo": nro_val,
                        "nombre_modulo": limpiar_texto(nombre_mod),
                        "horas_modulo": hrs_val,
                        "pago_docente_base": pago_val,
                    })

    # ──────────────────────────────────────────────────────────────────────────
    # 2. EJECUTADO VS META.xlsx
    # ──────────────────────────────────────────────────────────────────────────
    ruta_meta = os.path.join(dir_datos, "EJECUTADO VS META.xlsx")
    if os.path.exists(ruta_meta):
        print(f"Leyendo '{ruta_meta}'...")
        wb = openpyxl.load_workbook(ruta_meta, data_only=True)
        if "EJECUTADO" in wb.sheetnames:
            ws = wb["EJECUTADO"]
            # Columnas clave:
            # 0: CATEGORIA, 1: MES, 2: MES NUMERICO,
            # 4: META POR CATEGORIA 2024, 5: EJECUTADO 2024,
            # 7: META POR CATEGORIA 2025, 6: EJECUTADO 2025,
            # 9: META POR CATEGORIA 2026, 8: EJECUTADO 2026
            # 12: EBITDA META 2024, 13: EBITDA EJECUTADO 2024, 14: EBITDA META 2025
            for row in list(ws.iter_rows(values_only=True))[1:]:
                if len(row) < 10:
                    continue
                cat = limpiar_texto(row[0])
                mes_num = row[2]
                if not cat or not mes_num:
                    continue
                try:
                    m_val = int(float(mes_num))
                except (ValueError, TypeError):
                    continue

                if not (1 <= m_val <= 12):
                    continue

                # Parsear 2024, 2025, 2026
                def fnum(val):
                    if val is None or val == "":
                        return 0.0
                    try:
                        return float(str(val).replace(",", "."))
                    except ValueError:
                        return 0.0

                meta_24 = fnum(row[4])
                ejec_24 = fnum(row[5])
                meta_25 = fnum(row[7])
                ejec_25 = fnum(row[6])
                meta_26 = fnum(row[9])
                ejec_26 = fnum(row[8])

                ebitda_meta_24 = fnum(row[12]) if len(row) > 12 else 0.0
                ebitda_ejec_24 = fnum(row[13]) if len(row) > 13 else 0.0
                ebitda_meta_25 = fnum(row[14]) if len(row) > 14 else 0.0

                catalogos["metas"].append({
                    "gestion": 2024, "mes": m_val, "categoria": cat,
                    "meta": meta_24, "ejecutado": ejec_24,
                    "ebitda_meta": ebitda_meta_24, "ebitda_ejec": ebitda_ejec_24
                })
                catalogos["metas"].append({
                    "gestion": 2025, "mes": m_val, "categoria": cat,
                    "meta": meta_25, "ejecutado": ejec_25,
                    "ebitda_meta": ebitda_meta_25, "ebitda_ejec": 0.0
                })
                catalogos["metas"].append({
                    "gestion": 2026, "mes": m_val, "categoria": cat,
                    "meta": meta_26, "ejecutado": ejec_26,
                    "ebitda_meta": 0.0, "ebitda_ejec": 0.0
                })

    return catalogos


def generar_script_sql(catalogos: dict, ruta_salida: str):
    """Genera el script SQL de inicialización con sintaxis compatible ON CONFLICT."""
    lineas = [
        "-- =============================================================================",
        "-- SEMILLA DE CATÁLOGOS DE REFERENCIA INSTITUCIONAL 3NF",
        "-- Base de Datos: giragroup_catalog_db",
        "-- =============================================================================",
        "",
    ]

    # 1. Escuelas
    lineas.append("-- 1. Escuelas")
    for esc in sorted(catalogos["escuelas"]):
        lineas.append(
            f"INSERT INTO public.cat_escuela (nombre_escuela) VALUES ({escapar_sql(esc)}) "
            f"ON CONFLICT (nombre_escuela) DO NOTHING;"
        )
    lineas.append("")

    # 2. Tipos de Programa
    lineas.append("-- 2. Tipos de Programa")
    for tp in sorted(catalogos["tipos_programa"]):
        lineas.append(
            f"INSERT INTO public.cat_tipo_programa (nombre_tipo) VALUES ({escapar_sql(tp)}) "
            f"ON CONFLICT (nombre_tipo) DO NOTHING;"
        )
    lineas.append("")

    # 3. Modalidades
    lineas.append("-- 3. Modalidades")
    for mod in sorted(catalogos["modalidades"]):
        lineas.append(
            f"INSERT INTO public.cat_modalidad (nombre_modalidad) VALUES ({escapar_sql(mod)}) "
            f"ON CONFLICT (nombre_modalidad) DO NOTHING;"
        )
    lineas.append("")

    # 4. Estados Académicos
    lineas.append("-- 4. Criterios y Estados")
    for ea in sorted(catalogos["estados_academicos"]):
        lineas.append(
            f"INSERT INTO public.cat_estado_academico (nombre_estado) VALUES ({escapar_sql(ea)}) "
            f"ON CONFLICT (nombre_estado) DO NOTHING;"
        )
    for er in sorted(catalogos["estados_reprobados"]):
        lineas.append(
            f"INSERT INTO public.cat_estado_reprobado (nombre_detalle) VALUES ({escapar_sql(er)}) "
            f"ON CONFLICT (nombre_detalle) DO NOTHING;"
        )
    for ea in sorted(catalogos["estados_arca"]):
        lineas.append(
            f"INSERT INTO public.cat_estado_arca (nombre_arca) VALUES ({escapar_sql(ea)}) "
            f"ON CONFLICT (nombre_arca) DO NOTHING;"
        )
    lineas.append("")

    # 5. Clasificación de Cartera
    lineas.append("-- 5. Clasificación de Carteras de Cobranza")
    for car in catalogos["carteras"]:
        lineas.append(
            f"INSERT INTO public.cat_tipo_cartera (nombre_cartera, dias_mora_min, dias_mora_max, descripcion) "
            f"VALUES ({escapar_sql(car['nombre'])}, {car['min']}, {car['max']}, {escapar_sql(car['desc'])}) "
            f"ON CONFLICT (nombre_cartera) DO UPDATE SET "
            f"dias_mora_min = EXCLUDED.dias_mora_min, dias_mora_max = EXCLUDED.dias_mora_max, descripcion = EXCLUDED.descripcion;"
        )
    lineas.append("")

    # 6. Docentes
    lineas.append("-- 6. Docentes Canónicos")
    for doc in catalogos["docentes"].values():
        lineas.append(
            f"INSERT INTO public.cat_docente (nombre_completo, email, celular, grado_academico, tipo_docente) "
            f"VALUES ({escapar_sql(doc['nombre_completo'])}, {escapar_sql(doc['email'])}, "
            f"{escapar_sql(doc['celular'])}, {escapar_sql(doc['grado_academico'])}, {escapar_sql(doc['tipo_docente'])}) "
            f"ON CONFLICT (nombre_completo) DO UPDATE SET "
            f"email = COALESCE(EXCLUDED.email, public.cat_docente.email), "
            f"celular = COALESCE(EXCLUDED.celular, public.cat_docente.celular), "
            f"grado_academico = COALESCE(EXCLUDED.grado_academico, public.cat_docente.grado_academico);"
        )
    lineas.append("")

    # 7. Programas (asociando IDs de escuela, tipo, modalidad mediante subconsultas)
    lineas.append("-- 7. Programas (con versión y llaves foráneas resueltas)")
    for prg in catalogos["programas"].values():
        sql = (
            f"INSERT INTO public.cat_programa (pos_code, nombre_programa, version, nombre_completo_raw, "
            f"id_escuela, id_tipo_programa, id_modalidad, inversion_base, punto_equilibrio, horas_totales) "
            f"VALUES ("
            f"{escapar_sql(prg['pos_code'])}, {escapar_sql(prg['nombre_programa'])}, {prg['version']}, {escapar_sql(prg['nombre_completo_raw'])}, "
            f"(SELECT id_escuela FROM public.cat_escuela WHERE nombre_escuela = {escapar_sql(prg['escuela'])}), "
            f"(SELECT id_tipo_programa FROM public.cat_tipo_programa WHERE nombre_tipo = {escapar_sql(prg['tipo_programa'])}), "
            f"(SELECT id_modalidad FROM public.cat_modalidad WHERE nombre_modalidad = {escapar_sql(prg['modalidad'])}), "
            f"{prg['inversion_base']}, {prg['punto_equilibrio']}, {prg['horas_totales']}) "
            f"ON CONFLICT (pos_code) DO UPDATE SET "
            f"nombre_programa = EXCLUDED.nombre_programa, version = EXCLUDED.version, "
            f"id_escuela = EXCLUDED.id_escuela, id_tipo_programa = EXCLUDED.id_tipo_programa, "
            f"id_modalidad = EXCLUDED.id_modalidad, inversion_base = EXCLUDED.inversion_base, "
            f"punto_equilibrio = EXCLUDED.punto_equilibrio, horas_totales = EXCLUDED.horas_totales;"
        )
        lineas.append(sql)
    lineas.append("")

    # 8. Módulos
    lineas.append("-- 8. Módulos por Programa")
    for m in catalogos["modulos"]:
        sql = (
            f"INSERT INTO public.cat_modulo (id_programa, nro_modulo, nombre_modulo, horas_modulo, pago_docente_base) "
            f"SELECT id_programa, {m['nro_modulo']}, {escapar_sql(m['nombre_modulo'])}, {m['horas_modulo']}, {m['pago_docente_base']} "
            f"FROM public.cat_programa WHERE pos_code = {escapar_sql(m['pos_code'])} "
            f"ON CONFLICT (id_programa, nro_modulo) DO UPDATE SET "
            f"nombre_modulo = EXCLUDED.nombre_modulo, horas_modulo = EXCLUDED.horas_modulo, "
            f"pago_docente_base = EXCLUDED.pago_docente_base;"
        )
        lineas.append(sql)
    lineas.append("")

    # 9. Metas de Gestión
    lineas.append("-- 9. Metas de Gestión Presupuestaria")
    for meta in catalogos["metas"]:
        sql = (
            f"INSERT INTO public.ref_metas_gestion (gestion, mes, categoria, monto_meta_ingreso, monto_ejecutado, ebitda_meta_pct, ebitda_ejecutado_pct) "
            f"VALUES ({meta['gestion']}, {meta['mes']}, {escapar_sql(meta['categoria'])}, {meta['meta']}, {meta['ejecutado']}, "
            f"{meta['ebitda_meta']}, {meta['ebitda_ejec']}) "
            f"ON CONFLICT (gestion, mes, categoria) DO UPDATE SET "
            f"monto_meta_ingreso = EXCLUDED.monto_meta_ingreso, monto_ejecutado = EXCLUDED.monto_ejecutado, "
            f"ebitda_meta_pct = EXCLUDED.ebitda_meta_pct, ebitda_ejecutado_pct = EXCLUDED.ebitda_ejecutado_pct;"
        )
        lineas.append(sql)
    lineas.append("")

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    with open(ruta_salida, "w", encoding="utf-8") as f:
        f.write("\n".join(lineas))
    print(f"-> Script SQL de Catálogos generado: '{ruta_salida}'")


def sembrar_en_postgres_catalogos(ruta_sql: str) -> bool:
    """Ejecuta el script SQL en PostgreSQL si la base de datos de catálogos está disponible."""
    try:
        from sqlalchemy import create_engine, text
        from dotenv import load_dotenv
        load_dotenv()

        url = os.getenv("CATALOG_DATABASE_URL") or os.getenv("DATABASE_URL")
        if not url:
            return False

        if "giragroup_db" in url and not os.getenv("DOCKER_ENV"):
            url = url.replace("giragroup_db", "localhost")

        # Apuntar a giragroup_catalog_db
        if "giragroup_db" in url.split("/")[-1]:
            url = url.rsplit("/", 1)[0] + "/giragroup_catalog_db"

        engine = create_engine(url, connect_args={"connect_timeout": 3})
        with open(ruta_sql, "r", encoding="utf-8") as f:
            sql_script = f.read()

        with engine.begin() as conn:
            # Ejecutar sentencia por sentencia
            sentencias = [s.strip() for s in sql_script.split(";\n") if s.strip()]
            for s in sentencias:
                conn.execute(text(s))
        print("-> Inserción exitosa de catálogos en PostgreSQL (giragroup_catalog_db)!")
        return True
    except Exception as e:
        print(f"-> Nota: No se pudo conectar a PostgreSQL ({e}). El script SQL quedó listo para migración.")
        return False


def main():
    dir_datos = "uploads/data"
    print("Iniciando extracción de catálogos canónicos 3NF...")
    catalogos = extraer_datos_catalogos(dir_datos)

    print(f"Resumen de Catálogos extraídos:")
    print(f" - Escuelas: {len(catalogos['escuelas'])}")
    print(f" - Tipos de programa: {len(catalogos['tipos_programa'])}")
    print(f" - Modalidades: {len(catalogos['modalidades'])}")
    print(f" - Programas: {len(catalogos['programas'])}")
    print(f" - Docentes: {len(catalogos['docentes'])}")
    print(f" - Módulos: {len(catalogos['modulos'])}")
    print(f" - Criterios académicos: {len(catalogos['estados_academicos'])} estados, {len(catalogos['estados_reprobados'])} reprobados, {len(catalogos['estados_arca'])} ARCA")
    print(f" - Metas de gestión: {len(catalogos['metas'])}")

    ruta_sql = os.path.join(os.path.dirname(__file__), "seed_catalogos.sql")
    generar_script_sql(catalogos, ruta_sql)
    sembrar_en_postgres_catalogos(ruta_sql)


if __name__ == "__main__":
    main()
